"""文档解析：把上传文件转成纯文本 + 标题。

支持 .txt/.md/.csv/.docx/.xlsx/.pptx/.pdf；缺依赖时给出明确错误而不是静默失败。
"""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Dict, List

from .config import SUPPORTED_EXT
from .text import normalize_ws


class ParseError(Exception):
    pass


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError(f"无法识别文本编码: {path.name}")


def _parse_csv(path: Path) -> str:
    text = _read_text_file(path)
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return ""
    header = rows[0]
    lines = [" | ".join(header)]
    for row in rows[1:]:
        cells = [c.strip() for c in row]
        if any(cells):
            pairs = [f"{h}: {c}" for h, c in zip(header, cells) if c]
            lines.append("；".join(pairs) if pairs else " | ".join(cells))
    return "\n".join(lines)


def _parse_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ParseError("解析 .docx 需要 python-docx") from exc
    doc = Document(str(path))
    out: List[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading") or style.startswith("标题"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            out.append("#" * min(int(level), 6) + " " + t)
        else:
            out.append(t)
    for ti, table in enumerate(doc.tables, 1):
        out.append(f"## 表格 {ti}")
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        if rows:
            out.append(" | ".join(rows[0]))
            for row in rows[1:]:
                out.append(" | ".join(row))
    return "\n".join(out)


def _parse_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ParseError("解析 .xlsx 需要 openpyxl") from exc
    wb = load_workbook(str(path), data_only=True, read_only=True)
    out: List[str] = []
    for ws in wb.worksheets:
        out.append(f"# 工作表：{ws.title}")
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        for row in rows[1:]:
            cells = [str(c) if c is not None else "" for c in row]
            pairs = [f"{h}: {c}" for h, c in zip(header, cells) if c]
            if pairs:
                out.append("；".join(pairs))
    wb.close()
    return "\n".join(out)


def _parse_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise ParseError("解析 .pptx 需要 python-pptx") from exc
    prs = Presentation(str(path))
    out: List[str] = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"## 幻灯片 {i}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                out.append(shape.text.strip())
    return "\n".join(out)


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ParseError("解析 .pdf 需要 pypdf（pip install pypdf）") from exc
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            pages.append(f"## 第 {i} 页\n{t.strip()}")
    return "\n\n".join(pages)


_PARSERS = {
    ".txt": _read_text_file,
    ".md": _read_text_file,
    ".markdown": _read_text_file,
    ".csv": _parse_csv,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".pptx": _parse_pptx,
    ".pdf": _parse_pdf,
}


def parse_file(path: Path) -> Dict:
    """解析文件，返回 {title, text, ext, chars}。"""
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXT or ext not in _PARSERS:
        raise ParseError(f"不支持的文件类型: {ext}")
    text = _PARSERS[ext](path)
    text = normalize_ws(text)
    if not text:
        raise ParseError(f"文件内容为空: {path.name}")

    title = path.stem
    for line in text.splitlines():
        if line.startswith("#"):
            title = line.lstrip("#").strip() or title
            break
    return {"title": title, "text": text, "ext": ext, "chars": len(text), "name": path.name}